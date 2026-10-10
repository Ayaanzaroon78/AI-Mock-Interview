import io
import re
from pathlib import Path
from docx import Document
from pypdf import PdfReader


def extract_resume(file_storage, max_bytes=4 * 1024 * 1024, minimum=80):
    name = (file_storage.filename or "").lower()
    if not name or Path(name).suffix not in {".pdf", ".docx"}:
        raise ValueError("Upload a PDF or DOCX resume.")
    data = file_storage.read(max_bytes + 1)
    if not data:
        raise ValueError("The selected file is empty.")
    if len(data) > max_bytes:
        raise ValueError("File is larger than the 4 MB upload limit.")
    try:
        if name.endswith(".pdf"):
            if not data.startswith(b"%PDF-"):
                raise ValueError("This file does not appear to be a valid PDF.")
            reader = PdfReader(io.BytesIO(data), strict=True)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            if not data.startswith(b"PK\x03\x04"):
                raise ValueError("This file does not appear to be a valid DOCX document.")
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                text += "\n" + "\n".join(" ".join(c.text for c in row.cells) for row in table.rows)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("The document could not be read. Please try another PDF or DOCX file.") from exc
    text = re.sub(r"[\t\r ]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text).strip()
    if len(text) < minimum:
        if name.endswith(".pdf") and len(data) > 1000:
            raise ValueError("This PDF has too little selectable text. It may be scanned; OCR is not enabled.")
        raise ValueError("Not enough readable text was found in this resume.")
    return text[:30000]
